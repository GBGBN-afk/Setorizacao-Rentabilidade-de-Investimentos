#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tracker.py — versão 2
Rastreador de carteira: ações, ETFs (EUA), UCITS ETFs, fundos e renda fixa, em reais ou dólares,
com câmbio diário, inflação (IPCA), rentabilidade, volatilidade, setorização e gráficos.

FONTES DE DADOS
  Preços, dividendos e desdobramentos ..... Yahoo Finance, via yfinance (github.com/ranaroussi/yfinance)
  Setor de ações .......................... Yahoo Finance (campo "sector")
  Composição setorial de ETFs ............. 1º arquivo de posições do emissor (ex.: iShares "Download Holdings")
                                            2º Yahoo Finance (funds_data.sector_weightings)
                                            3º lâmina PDF, se a tabela de setores estiver em texto
                                            4º informada manualmente
  Câmbio (PTAX de fechamento, venda) ...... Banco Central — SGS série 1 (USD); API PTAX/Olinda (demais)
  IPCA mensal ............................. IBGE, via Banco Central — SGS série 433
  Projeção de IPCA (meses sem dado) ....... Boletim Focus — API Expectativas/Olinda (mediana mensal)
  Selic diária / CDI diário ............... Banco Central — SGS séries 11 e 12

Uso:  python tracker.py --help      |      guia completo no README.md
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import math
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from functools import lru_cache
from pathlib import Path

try:
    import numpy as np
    import pandas as pd
except ImportError:  # pragma: no cover
    sys.exit("Instale as dependências:  pip install -r requirements.txt")

VERSAO = 2
PASTA = Path(os.environ.get("TRACKER_PASTA", Path(__file__).resolve().parent))
ARQUIVO_PADRAO = PASTA / "carteira.json"
PASTA_CACHE = PASTA / "cache"
PASTA_GRAFICOS = PASTA / "graficos"


def hoje() -> dt.date:
    v = os.environ.get("TRACKER_HOJE")  # usado só nos testes
    return dt.date.fromisoformat(v) if v else dt.date.today()


def data(texto) -> dt.date:
    if isinstance(texto, dt.date):
        return texto
    t = str(texto).strip()
    if re.fullmatch(r"\d{2}/\d{2}/\d{4}", t):
        d, m, a = t.split("/")
        return dt.date(int(a), int(m), int(d))
    return dt.date.fromisoformat(t)


# =========================================================================================
# 1. Classificações
# =========================================================================================
SETORES_PT = {
    "technology": "Tecnologia", "financial_services": "Financeiro", "healthcare": "Saúde",
    "energy": "Energia", "consumer_cyclical": "Consumo Discricionário",
    "consumer_defensive": "Consumo Básico", "industrials": "Industrial", "basic_materials": "Materiais",
    "real_estate": "Imobiliário", "utilities": "Utilidades", "communication_services": "Comunicação",
    "renda_fixa": "Renda Fixa", "outros": "Outros / Não classificado",
}
_ALIASES = {
    "technology": "technology", "information_technology": "technology", "tech": "technology",
    "tecnologia": "technology", "tecnologia_da_informacao": "technology",
    "financial_services": "financial_services", "financials": "financial_services",
    "financial": "financial_services", "financeiro": "financial_services",
    "servicos_financeiros": "financial_services", "financas": "financial_services",
    "healthcare": "healthcare", "health_care": "healthcare", "saude": "healthcare",
    "energy": "energy", "energia": "energy", "petroleo_gas_e_biocombustiveis": "energy",
    "consumer_cyclical": "consumer_cyclical", "consumer_discretionary": "consumer_cyclical",
    "consumo_discricionario": "consumer_cyclical", "consumo_ciclico": "consumer_cyclical",
    "consumer_defensive": "consumer_defensive", "consumer_staples": "consumer_defensive",
    "consumo_basico": "consumer_defensive", "consumo_nao_ciclico": "consumer_defensive",
    "industrials": "industrials", "industrial": "industrials", "industria": "industrials",
    "bens_industriais": "industrials",
    "basic_materials": "basic_materials", "materials": "basic_materials",
    "materiais": "basic_materials", "materiais_basicos": "basic_materials",
    "real_estate": "real_estate", "realestate": "real_estate", "imobiliario": "real_estate",
    "imoveis": "real_estate",
    "utilities": "utilities", "utilidades": "utilities", "utilidade_publica": "utilities",
    "servicos_publicos": "utilities",
    "communication_services": "communication_services", "communication": "communication_services",
    "communications": "communication_services", "telecommunication_services": "communication_services",
    "comunicacao": "communication_services", "servicos_de_comunicacao": "communication_services",
    "telecomunicacoes": "communication_services",
    "renda_fixa": "renda_fixa",
    "outros": "outros", "other": "outros", "others": "outros", "cash": "outros", "caixa": "outros",
    "cash_and_or_derivatives": "outros", "cash_and_derivatives": "outros",
}
CLASSES_PLURAL = {"acao": "Ações", "etf": "ETFs (EUA)", "ucits": "UCITS ETFs", "fundo": "Fundos",
                  "renda_fixa": "Renda Fixa", "outro": "Outros"}
CLASSES = {"acao": "Ação", "etf": "ETF (EUA)", "ucits": "UCITS ETF", "fundo": "Fundo",
           "renda_fixa": "Renda Fixa", "outro": "Outro"}
QUOTE_TYPES = {"EQUITY": "acao", "ETF": "etf", "MUTUALFUND": "fundo"}
SUBUNIDADES = {"GBp": ("GBP", 100), "GBX": ("GBP", 100), "ZAc": ("ZAR", 100), "ILA": ("ILS", 100)}
SUFIXOS_EUROPA = {"L", "DE", "AS", "MI", "PA", "SW", "IR", "F", "BR", "MC", "VI", "CO", "ST", "HE", "OL"}
INDEXADORES = {"pre": "Prefixado", "cdi": "CDI", "selic": "Selic", "ipca": "IPCA+"}


def _slug(texto):
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", t.lower()).strip("_")


def normalizar_setor(texto):
    return _ALIASES.get(_slug(texto)) if texto else None


def normalizar_pesos(pesos_brutos):
    """{nome_setor: peso em % ou fração} -> ({chave: fração somando 1}, [não reconhecidos])."""
    out, desconhecidos = {}, []
    for nome, peso in pesos_brutos.items():
        try:
            peso = float(peso)
        except (TypeError, ValueError):
            continue
        if peso <= 0:
            continue
        chave = normalizar_setor(nome)
        if chave is None:
            desconhecidos.append(str(nome))
            chave = "outros"
        out[chave] = out.get(chave, 0.0) + peso
    total = sum(out.values())
    if total == 0:
        return {}, desconhecidos
    escala = 100.0 if total > 1.5 else 1.0
    out = {k: v / escala for k, v in out.items()}
    soma = sum(out.values())
    if soma < 0.999:
        out["outros"] = out.get("outros", 0.0) + (1 - soma)
    elif soma > 1.001:
        out = {k: v / soma for k, v in out.items()}
    return out, desconhecidos


def numero(texto):
    """Converte '1,234.56', '1.234,56', '13.05%', '-0,4' em float."""
    t = str(texto).strip().replace("%", "").replace("\u00a0", "").replace(" ", "").replace('"', "")
    if t in ("", "-", "--"):
        raise ValueError(texto)
    if "," in t and "." in t:
        t = t.replace(",", "") if t.rfind(".") > t.rfind(",") else t.replace(".", "").replace(",", ".")
    elif "," in t:
        t = t.replace(",", ".")
    return float(t)


# =========================================================================================
# 2. Calendário de dias úteis (feriados nacionais — calendário ANBIMA)
# =========================================================================================
def _pascoa(ano):
    a, b, c = ano % 19, ano // 100, ano % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = ((h + l - 7 * m + 114) % 31) + 1
    return dt.date(ano, mes, dia)


@lru_cache(maxsize=None)
def feriados(ano):
    p = _pascoa(ano)
    fixos = [(1, 1), (4, 21), (5, 1), (9, 7), (10, 12), (11, 2), (11, 15), (12, 25)]
    if ano >= 2024:
        fixos.append((11, 20))  # Consciência Negra, feriado nacional desde a Lei 14.759/2023
    s = {dt.date(ano, m, d) for m, d in fixos}
    s |= {p - dt.timedelta(days=48), p - dt.timedelta(days=47), p - dt.timedelta(days=2),
          p + dt.timedelta(days=60)}  # Carnaval (seg e ter), Sexta-feira Santa, Corpus Christi
    return frozenset(s)


def dia_util(d):
    return d.weekday() < 5 and d not in feriados(d.year)


def dias_uteis(d0, d1):
    """Dias úteis em [d0, d1)."""
    return sum(1 for i in range((d1 - d0).days) if dia_util(d0 + dt.timedelta(days=i)))


# =========================================================================================
# 3. Leitura de arquivos de composição (pura, testável)
# =========================================================================================
def ler_posicoes(conteudo: bytes, nome_arquivo: str):
    """Lê arquivo de posições (holdings) de um ETF e soma o peso por setor.
    Aceita CSV (iShares EUA/UK e outros emissores) ou XLSX. Retorna (pesos, data_ref, n_posicoes)."""
    if nome_arquivo.lower().endswith((".xlsx", ".xlsm")):
        df = pd.read_excel(io.BytesIO(conteudo), header=None, dtype=str)
        linhas = df.fillna("").values.tolist()
    else:
        texto = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                texto = conteudo.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        amostra = texto[:4000]
        sep = ";" if amostra.count(";") > amostra.count(",") else ","
        linhas = list(csv.reader(io.StringIO(texto), delimiter=sep))
    cab, col_setor, col_peso = None, None, None
    for i, linha in enumerate(linhas[:60]):
        norm = [_slug(x) for x in linha]
        setor = [j for j, x in enumerate(norm) if x in ("sector", "setor", "gics_sector", "sector_classification")]
        peso = [j for j, x in enumerate(norm) if x.startswith("weight") or x.startswith("peso")
                or x in ("pct_of_fund", "of_net_assets", "percent_of_fund", "weighting")]
        if setor and peso:
            cab, col_setor, col_peso = i, setor[0], peso[0]
            break
    if cab is None:
        raise ValueError("não encontrei as colunas de Setor e Peso (Weight) no arquivo")
    data_ref = None
    for linha in linhas[:cab]:
        junto = " ".join(str(x) for x in linha)
        m = re.search(r"(\d{1,2}[/\-. ]\w{3,9}[/\-. ]\d{4}|\d{4}-\d{2}-\d{2}|\d{2}/\d{2}/\d{4})", junto)
        if m and re.search(r"as of|holdings|posi|data", junto, re.I):
            data_ref = m.group(1)
            break
    pesos, n = {}, 0
    for linha in linhas[cab + 1:]:
        if len(linha) <= max(col_setor, col_peso) or not any(str(x).strip() for x in linha):
            if n:
                break
            continue
        try:
            p = numero(linha[col_peso])
        except ValueError:
            continue
        setor = str(linha[col_setor]).strip() or "outros"
        pesos[setor] = pesos.get(setor, 0.0) + p
        n += 1
    if not n:
        raise ValueError("nenhuma posição com peso numérico encontrada")
    return pesos, data_ref, n


def setores_de_texto_pdf(texto):
    padrao = re.compile(r"^\s*([A-Za-zÀ-ÿ&/,.\- ]{3,60}?)\s*[:\-–]?\s+(-?\d{1,3}(?:[.,]\d{1,3})?)\s*%?"
                        r"(?:\s+[\d.,%\-\s]*)?$")
    achados, vistos = {}, set()
    for linha in texto.splitlines():
        m = padrao.match(linha)
        if not m:
            continue
        nome, val = m.group(1).strip(), float(m.group(2).replace(",", "."))
        chave = normalizar_setor(nome)
        if chave and chave not in ("outros", "renda_fixa") and chave not in vistos and 0 < val <= 100:
            achados[nome] = val
            vistos.add(chave)
    return achados


def url_ishares_csv(url_pagina: str, ticker: str):
    """Monta o link do CSV de posições a partir da página do produto na iShares."""
    base = url_pagina.split("?")[0].split("#")[0].rstrip("/")
    if "ishares.com" not in base or "/products/" not in base:
        raise ValueError("use o link da página do produto, ex.: https://www.ishares.com/uk/individual/en/"
                         "products/297452/ishares-edge-msci-em-value-factor-ucits-etf")
    codigo = "1467271812596" if "/us/" in base else "1506575576011"
    nome = re.sub(r"[^A-Za-z0-9]", "", ticker.split(".")[0]) or "fund"
    return f"{base}/{codigo}.ajax?fileType=csv&fileName={nome}_holdings&dataType=fund"


# ---- parsers de respostas das APIs (puros, testáveis) -----------------------------------
def parse_sgs(lista):
    if not lista:
        return pd.Series(dtype=float)
    idx = pd.to_datetime([x["data"] for x in lista], format="%d/%m/%Y")
    vals = [float(str(x["valor"]).replace(",", ".")) for x in lista]
    return pd.Series(vals, index=idx, dtype=float).sort_index()


def parse_ptax_olinda(resposta):
    linhas = [v for v in resposta.get("value", []) if v.get("tipoBoletim", "Fechamento") == "Fechamento"]
    if not linhas:
        return pd.Series(dtype=float)
    idx = pd.to_datetime([v["dataHoraCotacao"][:10] for v in linhas])
    s = pd.Series([float(v["cotacaoVenda"]) for v in linhas], index=idx, dtype=float)
    return s[~s.index.duplicated(keep="last")].sort_index()


def parse_focus_mensal(resposta):
    """Mediana mais recente de IPCA mensal por mês de referência -> {'AAAA-MM': % no mês}."""
    out, vistos = {}, set()
    for v in sorted(resposta.get("value", []), key=lambda x: x["Data"], reverse=True):
        mm, aaaa = v["DataReferencia"].split("/")
        chave = f"{aaaa}-{mm}"
        if chave not in vistos and v.get("Mediana") is not None:
            out[chave] = float(v["Mediana"])
            vistos.add(chave)
    return out


# =========================================================================================
# 4. Fontes de dados reais (com cache em disco)
# =========================================================================================
class Cache:
    def __init__(self, pasta: Path):
        self.pasta = pasta

    def ler(self, chave):
        f = self.pasta / f"{_slug(chave)}.json"
        if f.exists():
            try:
                return json.loads(f.read_text(encoding="utf-8"))
            except json.JSONDecodeError:
                return None
        return None

    def gravar(self, chave, obj):
        self.pasta.mkdir(parents=True, exist_ok=True)
        (self.pasta / f"{_slug(chave)}.json").write_text(json.dumps(obj), encoding="utf-8")


def _serie_para_json(s):
    return {d.strftime("%Y-%m-%d"): (None if pd.isna(v) else float(v)) for d, v in s.items()}


def _serie_de_json(d):
    if not d:
        return pd.Series(dtype=float)
    s = pd.Series(d, dtype=float)
    s.index = pd.to_datetime(s.index)
    return s.sort_index()


class FontesReais:
    UA = {"User-Agent": "Mozilla/5.0 (tracker.py; uso pessoal)"}

    def __init__(self, pasta_cache=PASTA_CACHE, forcar=False):
        self.cache = Cache(pasta_cache)
        self.forcar = forcar
        self.avisos = []
        self._yf = None
        self._renovados = set()  # com forcar=True, cada série é baixada uma única vez por execução

    def _usar_cache(self, chave, c, ini=None):
        if not c or c.get("baixado") != hoje().isoformat():
            return False
        if ini is not None and c.get("ini", "9999") > ini.isoformat():
            return False
        return not self.forcar or chave in self._renovados

    # -- utilitários --------------------------------------------------------------------
    def yf(self):
        if self._yf is None:
            try:
                import yfinance
            except ImportError:
                sys.exit("Instale as dependências:  pip install -r requirements.txt")
            self._yf = yfinance
        return self._yf

    def baixar(self, url, timeout=30):
        req = urllib.request.Request(url, headers=self.UA)
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.read()

    def _json(self, url):
        return json.loads(self.baixar(url).decode("utf-8"))

    def _em_cache(self, chave, ini, fim, baixar):
        """Série temporal com cache diário. baixar(ini, fim) -> pd.Series."""
        c = self.cache.ler(chave)
        if self._usar_cache(chave, c, ini):
            return _serie_de_json(c["dados"])
        a = min(ini, dt.date.fromisoformat(c["ini"])) if c else ini
        try:
            s = baixar(a, fim)
        except Exception as e:
            if c:
                self.avisos.append(f"{chave}: sem conexão ({e.__class__.__name__}); usando dados salvos de {c['baixado']}")
                return _serie_de_json(c["dados"])
            raise
        self.cache.gravar(chave, {"baixado": hoje().isoformat(), "ini": a.isoformat(), "dados": _serie_para_json(s)})
        self._renovados.add(chave)
        return s

    # -- Banco Central ------------------------------------------------------------------
    def sgs(self, codigo, ini, fim):
        def baixar(a, b):
            partes = []
            while a <= b:  # a API limita consultas de séries diárias a 10 anos
                c = min(b, a + dt.timedelta(days=3600))
                url = (f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{codigo}/dados?formato=json"
                       f"&dataInicial={a:%d/%m/%Y}&dataFinal={c:%d/%m/%Y}")
                try:
                    partes.append(parse_sgs(self._json(url)))
                except urllib.error.HTTPError as e:
                    if e.code != 404:  # 404 = sem dados no intervalo
                        raise
                a = c + dt.timedelta(days=1)
            partes = [p for p in partes if len(p)]
            return pd.concat(partes).sort_index() if partes else pd.Series(dtype=float)
        return self._em_cache(f"sgs_{codigo}", ini, fim, baixar)

    def ptax(self, moeda, ini, fim):
        moeda = moeda.upper()

        def baixar(a, b):
            if moeda == "USD":
                return self.sgs(1, a, b)
            filtro = urllib.parse.quote("tipoBoletim eq 'Fechamento'")
            url = ("https://olinda.bcb.gov.br/olinda/servico/PTAX/versao/v1/odata/"
                   "CotacaoMoedaPeriodo(moeda=@moeda,dataInicial=@dataInicial,dataFinalCotacao=@dataFinalCotacao)"
                   f"?@moeda=%27{moeda}%27&@dataInicial=%27{a:%m-%d-%Y}%27&@dataFinalCotacao=%27{b:%m-%d-%Y}%27"
                   f"&$filter={filtro}&$format=json")
            try:
                s = parse_ptax_olinda(self._json(url))
                if len(s):
                    return s
            except Exception:
                pass
            self.avisos.append(f"PTAX {moeda} indisponível; usando cotação de mercado do Yahoo ({moeda}BRL=X)")
            h = self.yf().Ticker(f"{moeda}BRL=X").history(start=a, end=b + dt.timedelta(days=1))
            s = h["Close"].copy()
            s.index = s.index.tz_localize(None).normalize()
            return s
        return self._em_cache(f"ptax_{moeda}", ini, fim, baixar)

    def focus_ipca_mensal(self):
        c = self.cache.ler("focus_ipca")
        if self._usar_cache("focus_ipca", c):
            return c["dados"]
        filtro = urllib.parse.quote("Indicador eq 'IPCA' and baseCalculo eq 0")
        url = ("https://olinda.bcb.gov.br/olinda/servico/Expectativas/versao/v1/odata/ExpectativaMercadoMensais"
               f"?$filter={filtro}&$orderby={urllib.parse.quote('Data desc')}&$top=300&$format=json"
               "&$select=Data,DataReferencia,Mediana")
        try:
            dados = parse_focus_mensal(self._json(url))
            self.cache.gravar("focus_ipca", {"baixado": hoje().isoformat(), "dados": dados})
            self._renovados.add("focus_ipca")
            return dados
        except Exception as e:
            if c:
                self.avisos.append(f"Focus: sem conexão; usando projeções salvas de {c['baixado']}")
                return c["dados"]
            self.avisos.append(f"Focus indisponível ({e.__class__.__name__}); meses sem IPCA usarão média de 12 meses")
            return {}

    # -- Yahoo Finance ------------------------------------------------------------------
    def precos(self, ticker, ini, fim):
        """DataFrame diário com Close, Dividends, Splits (moeda já ajustada) e attrs['moeda']."""
        chave = f"px_{ticker}"
        c = self.cache.ler(chave)
        if self._usar_cache(chave, c, ini):
            return self._px_do_cache(c)
        a = min(ini, dt.date.fromisoformat(c["ini"])) if c else ini
        try:
            t = self.yf().Ticker(ticker)
            h = t.history(start=a, end=fim + dt.timedelta(days=1), auto_adjust=False, actions=True)
            if h is None or h.empty:
                raise RuntimeError(f"Yahoo não retornou cotações para {ticker}")
            moeda = (getattr(t, "history_metadata", None) or {}).get("currency")
            if not moeda:
                try:
                    moeda = t.fast_info["currency"]
                except Exception:
                    moeda = "USD"
            fator = 1.0
            if moeda in SUBUNIDADES:
                moeda, fator = SUBUNIDADES[moeda]
            df = pd.DataFrame({"Close": h["Close"] / fator,
                               "Dividends": h.get("Dividends", 0.0) / fator,
                               "Splits": h.get("Stock Splits", 0.0)})
            df.index = df.index.tz_localize(None).normalize()
            df = df[~df.index.duplicated(keep="last")]
        except Exception as e:
            if c:
                self.avisos.append(f"{ticker}: sem conexão com Yahoo ({e.__class__.__name__}); usando cotações salvas")
                return self._px_do_cache(c)
            raise
        self.cache.gravar(chave, {"baixado": hoje().isoformat(), "ini": a.isoformat(), "moeda": moeda,
                                  "dados": {d.strftime("%Y-%m-%d"): [float(r.Close), float(r.Dividends),
                                                                      float(r.Splits)]
                                            for d, r in df.iterrows()}})
        self._renovados.add(chave)
        df.columns = ["Close", "Dividends", "Splits"]
        df.attrs["moeda"] = moeda
        return df

    @staticmethod
    def _px_do_cache(c):
        df = pd.DataFrame.from_dict(c["dados"], orient="index", columns=["Close", "Dividends", "Splits"]).astype(float)
        df.index = pd.to_datetime(df.index)
        df = df.sort_index()
        df.attrs["moeda"] = c["moeda"]
        return df

    def info(self, ticker):
        t = self.yf().Ticker(ticker)
        try:
            i = t.info or {}
        except Exception:
            i = {}
        moeda = i.get("currency")
        if moeda in SUBUNIDADES:
            moeda = SUBUNIDADES[moeda][0]
        return {"nome": i.get("longName") or i.get("shortName") or ticker,
                "quote_type": (i.get("quoteType") or "").upper(), "moeda": moeda, "setor": i.get("sector")}

    def pesos_fundo(self, ticker):
        try:
            p = self.yf().Ticker(ticker).funds_data.sector_weightings
        except Exception:
            return None
        return dict(p) if p else None


# ---- dados cadastrados em JSON antigo (v1) viram v2 ------------------------------------
def migrar_v1(c):
    novo = {"versao": VERSAO, "ativos": {}, "movimentos": [], "valores": {}, "config": {}}
    classe_de = {"acao": "acao", "etf": "etf", "fundo": "fundo", "outro": "outro"}
    for t, a in c.get("ativos", {}).items():
        novo["ativos"][t] = {
            "nome": a.get("nome", t), "classe": classe_de.get(a.get("tipo"), "outro"), "moeda": a.get("moeda", "BRL"),
            "fonte_preco": "yahoo" if a.get("modo") == "auto" else "manual",
            "setores": a.get("setores", {}), "fonte_setores": a.get("fonte_setores"),
            "data_setores": None,
            "retencao_dividendos": 0.30 if a.get("moeda") == "USD" and a.get("tipo") in ("acao", "etf") else 0.0,
        }
    for i, m in enumerate(c.get("movimentos", []), start=1):
        a = c["ativos"].get(m["ticker"], {})
        auto = a.get("modo") == "auto"
        q = m.get("quantidade")
        if auto and q:
            tipo = "compra" if m["tipo"] == "aporte" else "venda"
            novo["movimentos"].append({"id": i, "data": m["data"], "ativo": m["ticker"], "tipo": tipo,
                                       "quantidade": q, "preco": m["valor"] / q, "corretagem": 0.0,
                                       "cambio": m.get("cambio_brl"), "fonte_cambio": "PTAX (migrado da v1)"})
        else:
            novo["movimentos"].append({"id": i, "data": m["data"], "ativo": m["ticker"],
                                       "tipo": "aporte" if m["tipo"] == "aporte" else "resgate",
                                       "valor": m["valor"], "cambio": m.get("cambio_brl"),
                                       "fonte_cambio": "PTAX (migrado da v1)"})
    for h in c.get("historico", []):
        a = c["ativos"].get(h["ticker"], {})
        if a.get("modo") != "auto":
            novo["valores"].setdefault(h["ticker"], []).append({"data": h["data"], "valor": h["valor"]})
    return novo


def carregar(caminho: Path):
    if not caminho.exists():
        return {"versao": VERSAO, "ativos": {}, "movimentos": [], "valores": {}, "config": {}}
    c = json.loads(caminho.read_text(encoding="utf-8"))
    if c.get("versao") != VERSAO:
        backup = caminho.with_suffix(".v1.bak.json")
        backup.write_text(json.dumps(c, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"• Carteira no formato antigo convertida para v2 (cópia do original em {backup.name}).")
        c = migrar_v1(c)
    for k, v in (("valores", {}), ("config", {})):
        c.setdefault(k, v)
    return c


def salvar(c, caminho: Path):
    tmp = caminho.with_suffix(".tmp")
    tmp.write_text(json.dumps(c, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(caminho)


# =========================================================================================
# 5. Motor de avaliação diária
# =========================================================================================
class Motor:
    """Monta, para cada ativo, a série diária de valor e fluxos na moeda do ativo e em reais."""

    def __init__(self, carteira, fontes, ate=None):
        self.c, self.f = carteira, fontes
        self.ate = ate or hoje()
        self._series, self._fx, self._ipca_mensal = {}, {}, None
        self.fonte_ipca = {}
        self.avisos = []

    def movs(self, t):
        return sorted([m for m in self.c["movimentos"] if m["ativo"] == t and data(m["data"]) <= self.ate],
                      key=lambda m: (m["data"], m["id"]))

    def inicio_carteira(self):
        ds = [data(m["data"]) for m in self.c["movimentos"] if data(m["data"]) <= self.ate]
        return min(ds) if ds else self.ate

    def _idx(self, ini):
        return pd.date_range(pd.Timestamp(ini), pd.Timestamp(self.ate), freq="D")

    # -- câmbio ---------------------------------------------------------------------------
    def fx(self, moeda, idx):
        moeda = (moeda or "BRL").upper()
        if moeda == "BRL":
            return pd.Series(1.0, index=idx)
        if moeda not in self._fx:
            ini = self.inicio_carteira() - dt.timedelta(days=10)
            s = self.f.ptax(moeda, ini, self.ate)
            if not len(s):
                raise RuntimeError(f"sem cotação de câmbio para {moeda}")
            self._fx[moeda] = s
        s = self._fx[moeda]
        return s.reindex(s.index.union(idx)).ffill().bfill().reindex(idx)

    def cambio_do_movimento(self, m, moeda, fx_serie):
        if (moeda or "BRL").upper() == "BRL":
            return 1.0
        return float(m["cambio"]) if m.get("cambio") else float(fx_serie.loc[pd.Timestamp(m["data"])])

    # -- IPCA -----------------------------------------------------------------------------
    def ipca_mensal(self, ini, fim):
        """{(ano, mês): taxa decimal} de ini a fim, com a origem de cada mês em self.fonte_ipca."""
        if self._ipca_mensal is None:
            base = min(ini, self.inicio_carteira())
            oficial = self.f.sgs(433, dt.date(base.year - 1, 1, 1), max(fim, self.ate))
            self._oficial = {(d.year, d.month): v / 100 for d, v in oficial.items()}
            self._focus = self.f.focus_ipca_mensal()
            self._ipca_mensal = {}
        proj = self.c.get("config", {}).get("ipca_projecoes", {})
        ult12 = [v for k, v in sorted(self._oficial.items())][-12:]
        media = (np.prod([1 + v for v in ult12]) ** (1 / len(ult12)) - 1) if ult12 else 0.004
        a, m = ini.year, ini.month
        while (a, m) <= (fim.year, fim.month):
            if (a, m) not in self._ipca_mensal:
                chave = f"{a}-{m:02d}"
                if (a, m) in self._oficial:
                    v, fonte = self._oficial[(a, m)], "IBGE"
                elif chave in proj:
                    v, fonte = proj[chave] / 100, "sua projeção"
                elif chave in self._focus:
                    v, fonte = self._focus[chave] / 100, "Focus (mediana)"
                else:
                    v, fonte = media, "média 12 meses"
                self._ipca_mensal[(a, m)] = v
                self.fonte_ipca[chave] = fonte
            a, m = (a + 1, 1) if m == 12 else (a, m + 1)
        return self._ipca_mensal

    def indice_ipca(self, idx):
        """Índice diário do IPCA, pró-rata por dias corridos dentro de cada mês."""
        ini, fim = idx[0].date(), idx[-1].date()
        taxas = self.ipca_mensal(dt.date(ini.year, ini.month, 1), fim)
        base, acumulado = {}, 1.0
        a, m = ini.year, ini.month
        while (a, m) <= (fim.year, fim.month):
            base[(a, m)] = acumulado
            acumulado *= 1 + taxas[(a, m)]
            a, m = (a + 1, 1) if m == 12 else (a, m + 1)
        vals = []
        for d in idx:
            dias_mes = pd.Timestamp(d).days_in_month
            vals.append(base[(d.year, d.month)] * (1 + taxas[(d.year, d.month)]) ** ((d.day - 1) / dias_mes))
        return pd.Series(vals, index=idx)

    # -- séries por ativo -----------------------------------------------------------------
    def serie(self, t):
        if t in self._series:
            return self._series[t]
        a = self.c["ativos"][t]
        movs = self.movs(t)
        if not movs:
            self._series[t] = None
            return None
        if a["fonte_preco"] == "yahoo":
            df = self._serie_mercado(t, a, movs)
        elif a["fonte_preco"] == "renda_fixa":
            df = self._serie_rf(t, a, movs)
        else:
            df = self._serie_manual(t, a, movs)
        fx = self.fx(a["moeda"], df.index)
        df["fx"] = fx
        df["v_brl"] = df["v_nat"] * fx
        cf_brl = pd.Series(0.0, index=df.index)
        for m in movs:
            fluxo = m.get("_fluxo_nat", 0.0)
            if fluxo:
                cf_brl.loc[pd.Timestamp(m["data"])] += fluxo * self.cambio_do_movimento(m, a["moeda"], fx)
        df["cf_brl"] = cf_brl + df.get("cf_brl_extra", 0.0)
        self._series[t] = df
        return df

    def _serie_mercado(self, t, a, movs):
        ini = data(movs[0]["data"])
        idx = self._idx(ini)
        px = self.f.precos(t, ini - dt.timedelta(days=10), self.ate)
        if px.attrs.get("moeda") and px.attrs["moeda"] != a["moeda"]:
            self.avisos.append(f"{t}: Yahoo cota em {px.attrs['moeda']}, cadastro diz {a['moeda']} — usando {px.attrs['moeda']}")
            a["moeda"] = px.attrs["moeda"]
        close = px["Close"].reindex(px.index.union(idx)).ffill().bfill().reindex(idx)
        splits = px["Splits"][px["Splits"] > 0] if "Splits" in px else pd.Series(dtype=float)

        def fator_split(d):  # cotações do Yahoo já vêm ajustadas a desdobramentos posteriores
            return float(np.prod([r for ds, r in splits.items() if ds > pd.Timestamp(d)])) if len(splits) else 1.0

        dq = pd.Series(0.0, index=idx)
        cf = pd.Series(0.0, index=idx)
        dcaixa = pd.Series(0.0, index=idx)
        for m in movs:
            d = pd.Timestamp(m["data"])
            k = fator_split(d)
            q, p, corr = (m.get("quantidade") or 0) * k, (m.get("preco") or 0) / k, m.get("corretagem") or 0
            if m["tipo"] == "compra":
                dq.loc[d] += q
                m["_fluxo_nat"] = q * p + corr
            elif m["tipo"] == "venda":
                dq.loc[d] -= q
                m["_fluxo_nat"] = -(q * p - corr)
            elif m["tipo"] == "reinvestir":
                dq.loc[d] += q
                dcaixa.loc[d] -= q * p + corr
                m["_fluxo_nat"] = 0.0
            elif m["tipo"] == "retirar":
                dcaixa.loc[d] -= m["valor"]
                m["_fluxo_nat"] = -m["valor"]
            cf.loc[d] += m.get("_fluxo_nat", 0.0)
        qtd = dq.cumsum()
        ret = a.get("retencao_dividendos", 0.0)
        divs = px["Dividends"][px["Dividends"] > 0] if "Dividends" in px else pd.Series(dtype=float)
        qtd_antes = qtd.shift(1, fill_value=0.0)
        for d, v in divs.items():
            if d in idx:
                dcaixa.loc[d] += qtd_antes.loc[d] * v * (1 - ret)
        caixa, extra, saldo = [], pd.Series(0.0, index=idx), 0.0
        for d in idx:  # reinvestir mais do que os dividendos disponíveis = dinheiro novo
            saldo += dcaixa.loc[d]
            if saldo < -1e-9:
                extra.loc[d] += -saldo
                saldo = 0.0
            caixa.append(saldo)
        caixa = pd.Series(caixa, index=idx)
        df = pd.DataFrame({"qtd": qtd, "preco": close, "caixa_div": caixa}, index=idx)
        df["v_nat"] = qtd * close + caixa
        df["cf_nat"] = cf + extra
        if extra.sum() > 0:
            fx = self.fx(a["moeda"], idx)
            df["cf_brl_extra"] = extra * fx
        return df

    def _serie_manual(self, t, a, movs):
        idx = self._idx(data(movs[0]["data"]))
        cf = pd.Series(0.0, index=idx)
        for m in movs:
            v = m["valor"] if m["tipo"] == "aporte" else -m["valor"]
            m["_fluxo_nat"] = v
            cf.loc[pd.Timestamp(m["data"])] += v
        relatos = {pd.Timestamp(r["data"]): r["valor"] for r in self.c["valores"].get(t, [])
                   if data(r["data"]) <= self.ate}
        acum = cf.cumsum()
        vals, base_v, base_acum = [], None, 0.0
        for d in idx:
            if d in relatos:
                base_v, base_acum = relatos[d], acum.loc[d]
            vals.append(acum.loc[d] if base_v is None else base_v + acum.loc[d] - base_acum)
        return pd.DataFrame({"v_nat": vals, "cf_nat": cf}, index=idx)

    def fator_rf(self, rf, idx):
        """G(t): valor em t de 1 unidade aplicada no início do índice. F(d0,t) = G(t)/G(d0)."""
        ind = rf["indexador"]
        uteis = np.array([dia_util(d.date()) for d in idx])
        du_ate = pd.Series(np.concatenate([[0], np.cumsum(uteis)[:-1]]), index=idx)  # dias úteis em [início, t)
        if ind == "pre":
            if rf.get("base", 252) == 365:
                expo = pd.Series(np.arange(len(idx)) / 365.0, index=idx)
            else:
                expo = du_ate / 252.0
            return (1 + rf["taxa"]) ** expo
        if ind in ("cdi", "selic"):
            s = self.f.sgs(12 if ind == "cdi" else 11, idx[0].date() - dt.timedelta(days=10), self.ate)
            ultimo = s.index.max() if len(s) else pd.Timestamp(idx[0]) - pd.Timedelta(days=1)
            ult_taxa = float(s.iloc[-1]) if len(s) else 0.0
            pct, spread = rf.get("percentual", 1.0), rf.get("spread", 0.0)
            fd = []
            for d, util in zip(idx, uteis):
                if d in s.index:
                    r = float(s.loc[d])
                elif d > ultimo and util:  # taxa do dia ainda não publicada: repete a última
                    r = ult_taxa
                else:
                    fd.append(1.0)
                    continue
                fd.append((1 + r / 100 * pct) * (1 + spread) ** (1 / 252))
            fd = pd.Series(fd, index=idx)
            return fd.cumprod().shift(1, fill_value=1.0)
        if ind == "ipca":
            ipca = self.indice_ipca(idx)
            return ipca / ipca.iloc[0] * (1 + rf["taxa"]) ** (du_ate / 252.0)
        raise ValueError(f"indexador desconhecido: {ind}")

    def _serie_rf(self, t, a, movs):
        idx = self._idx(data(movs[0]["data"]))
        g = self.fator_rf(a["rf"], idx).to_numpy()
        lotes, cf = [], pd.Series(0.0, index=idx)
        for m in movs:
            pos = idx.get_loc(pd.Timestamp(m["data"]))
            if m["tipo"] == "aporte":
                arr = np.zeros(len(idx))
                arr[pos:] = m["valor"] * g[pos:] / g[pos]
                lotes.append({"data": m["data"], "principal": m["valor"], "valores": arr})
                m["_fluxo_nat"] = m["valor"]
            else:
                total = sum(l["valores"][pos] for l in lotes)
                if m["valor"] > total + 1e-6:
                    raise ValueError(f"{t}: resgate de {m['valor']:.2f} maior que o saldo {total:.2f}")
                f = m["valor"] / total if total else 0
                for l in lotes:
                    l["valores"][pos:] *= (1 - f)
                    l["principal"] *= (1 - f)
                m["_fluxo_nat"] = -m["valor"]
            cf.iloc[pos] += m["_fluxo_nat"]
        a["_lotes"] = lotes
        v = np.sum([l["valores"] for l in lotes], axis=0)
        return pd.DataFrame({"v_nat": v, "cf_nat": cf}, index=idx)

    def ir_estimado_rf(self, t):
        """IR regressivo sobre o rendimento de cada lote (não considera IOF dos primeiros 30 dias)."""
        a = self.c["ativos"][t]
        if a.get("rf", {}).get("isento") or self.serie(t) is None:
            return 0.0
        ir = 0.0
        for l in a.get("_lotes", []):
            ganho = l["valores"][-1] - l["principal"]
            dias = (self.ate - data(l["data"])).days
            aliq = 0.225 if dias <= 180 else 0.20 if dias <= 360 else 0.175 if dias <= 720 else 0.15
            ir += max(ganho, 0) * aliq
        return ir


# =========================================================================================
# 6. Métricas
# =========================================================================================
def _alinhar(frames, idx):
    out = {}
    for t, df in frames.items():
        d = df.reindex(idx)
        for c in ("v_nat", "cf_nat", "v_brl", "cf_brl"):
            d[c] = d[c].fillna(0.0)
        d["fx"] = d["fx"].ffill().bfill()
        out[t] = d
    return out


def _retornos(num, den):
    return pd.Series(np.where(den > 1e-9, num / np.where(den > 1e-9, den, 1) - 1, 0.0), index=num.index)


def calcular(motor, tickers, desde=None, anualizar_min=365):
    """Métricas do conjunto de ativos no período (desde, motor.ate].
    Anualiza só períodos de 1 ano ou mais (padrão GIPS/CFA), salvo anualizar_min menor."""
    frames = {t: motor.serie(t) for t in tickers}
    frames = {t: f for t, f in frames.items() if f is not None}
    if not frames:
        return None
    ini_total = min(f.index[0] for f in frames.values())
    idx = pd.date_range(ini_total, pd.Timestamp(motor.ate), freq="D")
    fr = _alinhar(frames, idx)

    def partes(col_v, col_cf, d):
        v, cf = d[col_v], d[col_cf]
        return v - cf.clip(upper=0), v.shift(1, fill_value=0.0) + cf.clip(lower=0)

    num = sum(partes("v_brl", "cf_brl", d)[0] for d in fr.values())
    den = sum(partes("v_brl", "cf_brl", d)[1] for d in fr.values())
    r_com = _retornos(num, den)
    r_sem = pd.Series(0.0, index=idx)
    for d in fr.values():  # sem câmbio: retorno de cada ativo na própria moeda, pesado pela alocação em R$
        n_nat, d_nat = partes("v_nat", "cf_nat", d)
        peso = partes("v_brl", "cf_brl", d)[1] / den.where(den > 1e-9, np.nan)
        r_sem += (_retornos(n_nat, d_nat) * peso.fillna(0.0))

    inicio = pd.Timestamp(desde) if desde else idx[0]
    if inicio < idx[0]:
        inicio = idx[0]
    mask = (idx > inicio) if desde else (idx >= idx[0])
    dias = (idx[-1] - inicio).days
    tw_com = float(np.prod(1 + r_com[mask])) - 1
    tw_sem = float(np.prod(1 + r_sem[mask])) - 1
    ipca_idx = motor.indice_ipca(idx)
    ipca_per = float(ipca_idx.iloc[-1] / ipca_idx.loc[inicio]) - 1

    def anual(r):
        return (1 + r) ** (365 / dias) - 1 if dias >= max(anualizar_min, 1) else None

    def vol(r):
        nivel = (1 + r[mask]).cumprod()
        sem = nivel.resample("W-FRI").last().pct_change().dropna()
        return (float(sem.std(ddof=1) * math.sqrt(52)) if len(sem) >= 4 else None), len(sem)

    v_com, n_sem = vol(r_com)
    v_sem, _ = vol(r_sem)
    V = sum(d["v_brl"] for d in fr.values())
    CF = sum(d["cf_brl"] for d in fr.values())
    V0 = float(V.loc[inicio]) if desde else 0.0
    fluxos = CF[mask]
    resultado = float(V.iloc[-1]) - V0 - float(fluxos.sum())
    res_sem = 0.0
    for t, d in fr.items():
        nat_fluxos = d["cf_nat"][mask]
        v0n = float(d["v_nat"].loc[inicio]) if desde else 0.0
        if desde:
            fx_ref = float(d["fx"].loc[inicio])
        else:
            compras = d["cf_nat"] > 0
            fx_ref = float(d["cf_brl"][compras].sum() / d["cf_nat"][compras].sum()) if compras.any() else float(d["fx"].iloc[-1])
        res_sem += (float(d["v_nat"].iloc[-1]) - v0n - float(nat_fluxos.sum())) * fx_ref
    fator_fim = float(ipca_idx.iloc[-1])
    real = float(V.iloc[-1]) - V0 * fator_fim / float(ipca_idx.loc[inicio]) - float(
        (fluxos * (fator_fim / ipca_idx[mask])).sum())
    aportes = float(CF[mask].clip(lower=0).sum())
    return {
        "inicio": inicio.date(), "fim": idx[-1].date(), "dias": dias,
        "twr_com": tw_com, "twr_sem": tw_sem, "ipca": ipca_per,
        "real_com": (1 + tw_com) / (1 + ipca_per) - 1, "real_sem": (1 + tw_sem) / (1 + ipca_per) - 1,
        "anual_com": anual(tw_com), "anual_sem": anual(tw_sem),
        "anual_real_com": anual((1 + tw_com) / (1 + ipca_per) - 1),
        "anual_real_sem": anual((1 + tw_sem) / (1 + ipca_per) - 1),
        "vol_com": v_com, "vol_sem": v_sem, "semanas": n_sem,
        "anualizar_min": anualizar_min, "valor_brl": float(V.iloc[-1]), "aportes": aportes, "fluxo_liquido": float(fluxos.sum()),
        "resultado": resultado, "resultado_sem_cambio": res_sem, "efeito_cambio": resultado - res_sem,
        "resultado_real": real, "r_com": r_com, "r_sem": r_sem, "idx": idx, "ipca_idx": ipca_idx,
        "V": V, "CF": CF,
    }


def alocacao(motor, tickers, por="classe"):
    valores = {}
    for t in tickers:
        df = motor.serie(t)
        if df is None:
            continue
        v = float(df["v_brl"].iloc[-1])
        a = motor.c["ativos"][t]
        if por == "classe":
            valores[CLASSES[a["classe"]]] = valores.get(CLASSES[a["classe"]], 0.0) + v
        else:
            pesos = {"renda_fixa": 1.0} if a["classe"] == "renda_fixa" else (a.get("setores") or {"outros": 1.0})
            for k, p in pesos.items():
                valores[SETORES_PT[k]] = valores.get(SETORES_PT[k], 0.0) + v * p
    return dict(sorted(valores.items(), key=lambda x: -x[1]))


def retorno_mensal_setores(motor, tickers, cambio=True):
    """Retorno mensal de cada setor: média dos retornos dos ativos ponderada por (valor × peso do setor)."""
    frames = {t: motor.serie(t) for t in tickers}
    frames = {t: f for t, f in frames.items() if f is not None}
    if not frames:
        return pd.DataFrame()
    idx = pd.date_range(min(f.index[0] for f in frames.values()), pd.Timestamp(motor.ate), freq="D")
    fr = _alinhar(frames, idx)
    meses = idx.to_period("M")
    ret_ativo, peso_ativo = {}, {}
    for t, d in fr.items():
        cv, cc = ("v_brl", "cf_brl") if cambio else ("v_nat", "cf_nat")
        num = d[cv] - d[cc].clip(upper=0)
        den = d[cv].shift(1, fill_value=0.0) + d[cc].clip(lower=0)
        r = _retornos(num, den)
        base_brl = d["v_brl"].shift(1, fill_value=0.0) + d["cf_brl"].clip(lower=0)
        rm, wm = {}, {}
        for p in meses.unique():
            sel = (meses == p) & (den > 1e-9)
            if sel.any():
                rm[p] = float(np.prod(1 + r[sel])) - 1
                wm[p] = float(base_brl[sel].iloc[0])
        ret_ativo[t], peso_ativo[t] = rm, wm
    linhas = {}
    for p in meses.unique():
        por_setor = {}
        for t in fr:
            if p not in ret_ativo[t]:
                continue
            a = motor.c["ativos"][t]
            pesos = {"renda_fixa": 1.0} if a["classe"] == "renda_fixa" else (a.get("setores") or {"outros": 1.0})
            for k, e in pesos.items():
                w = e * peso_ativo[t][p]
                acc = por_setor.setdefault(SETORES_PT[k], [0.0, 0.0])
                acc[0] += w * ret_ativo[t][p]
                acc[1] += w
        linhas[str(p)] = {s: n / w for s, (n, w) in por_setor.items() if w > 0}
    df = pd.DataFrame(linhas).T
    carteira = calcular(motor, list(fr))
    rc = carteira["r_com"] if cambio else carteira["r_sem"]
    df["Carteira"] = [float(np.prod(1 + rc[meses == pd.Period(m)])) - 1 for m in df.index]
    parciais = set()
    if idx[0].day != 1:
        parciais.add(str(idx[0].to_period("M")))
    if not idx[-1].is_month_end:
        parciais.add(str(idx[-1].to_period("M")))
    df.attrs["parciais"] = parciais
    return df


# =========================================================================================
# 7. Formatação de saída
# =========================================================================================
def br(x, casas=2):
    s = f"{x:,.{casas}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, sinal=True):
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return "—"
    return (("+" if x >= 0 else "") if sinal else "") + br(x * 100, 2) + "%"


def rs(x):
    return "—" if x is None else ("-" if x < 0 else "") + "R$ " + br(abs(x))


def tabela(cab, linhas, direita=None):
    direita = set(range(1, len(cab))) if direita is None else set(direita)
    larg = [max(len(str(c)), *(len(str(l[i])) for l in linhas)) if linhas else len(str(c)) for i, c in enumerate(cab)]
    fmt = lambda row: "  ".join(str(v).rjust(larg[i]) if i in direita else str(v).ljust(larg[i]) for i, v in enumerate(row))
    print(fmt(cab))
    print("  ".join("─" * w for w in larg))
    for l in linhas:
        print(fmt(l))


# =========================================================================================
# 8. Gráficos
# =========================================================================================
MESES_PT = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]


def mes_pt(periodo):
    p = pd.Period(periodo)
    return f"{MESES_PT[p.month - 1]}/{p.year}"


def _plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    return plt


def grafico_pizza(valores, titulo, arquivo, minimo=0.01):
    plt = _plt()
    total = sum(valores.values())
    grandes = {k: v for k, v in valores.items() if v / total >= minimo}
    pequenos = {k: v for k, v in valores.items() if v / total < minimo}
    if len(pequenos) > 1:
        grandes[f"Demais ({len(pequenos)} itens, <{minimo:.0%} cada)"] = sum(pequenos.values())
    else:
        grandes.update(pequenos)
    valores = grandes
    cores = plt.get_cmap("tab20").colors
    fig, ax = plt.subplots(figsize=(9, 6))
    fatias, _ = ax.pie(list(valores.values()), startangle=90, counterclock=False,
                       colors=[cores[i % 20] for i in range(len(valores))],
                       wedgeprops={"linewidth": 1, "edgecolor": "white"})
    for f, (k, v) in zip(fatias, valores.items()):
        if v / total >= 0.03:
            ang = (f.theta2 + f.theta1) / 2
            x, y = 0.68 * math.cos(math.radians(ang)), 0.68 * math.sin(math.radians(ang))
            ax.text(x, y, f"{v / total:.0%}", ha="center", va="center", fontsize=10, color="white", weight="bold")
    ax.legend(fatias, [f"{k} — {br(v / total * 100, 1)}%" for k, v in valores.items()],
              loc="center left", bbox_to_anchor=(1, 0.5), frameon=False)
    ax.set_title(titulo, fontsize=13, weight="bold")
    ax.axis("equal")
    fig.tight_layout()
    fig.savefig(arquivo, dpi=150, bbox_inches="tight")
    plt.close(fig)


def grafico_linhas(df, titulo, arquivo, eixo_y="Rentabilidade no mês", acumulado=False, nota=None):
    plt = _plt()
    fig, ax = plt.subplots(figsize=(11, 6))
    dados = ((1 + df.fillna(0)).cumprod() - 1) if acumulado else df
    x = list(range(len(dados.index)))
    for col in dados.columns:
        serie = dados[col]
        estilo = {"linestyle": "--", "linewidth": 2.6, "color": "black"} if col == "Carteira" else {"linewidth": 1.8}
        ax.plot(x, serie.values * 100, marker="o", markersize=4, label=col, **estilo)
    ax.axhline(0, color="grey", linewidth=0.8)
    ax.set_xticks(x)
    parciais = getattr(df, "attrs", {}).get("parciais", set())
    ax.set_xticklabels([mes_pt(i) + ("*" if str(i) in parciais else "") for i in dados.index], rotation=45, ha="right")
    if parciais:
        nota = ((nota + "  ") if nota else "") + "* mês parcial (início da carteira ou mês corrente)."
    ax.set_ylabel(eixo_y + " (%)")
    ax.set_title(titulo, fontsize=13, weight="bold")
    ax.grid(alpha=0.3)
    ax.legend(loc="center left", bbox_to_anchor=(1, 0.5), frameon=False)
    if nota:
        fig.text(0.01, -0.02, nota, fontsize=8, color="grey", ha="left")
    fig.tight_layout()
    fig.savefig(arquivo, dpi=150, bbox_inches="tight")
    plt.close(fig)


def grafico_evolucao(m, titulo, arquivo):
    plt = _plt()
    fig, ax = plt.subplots(figsize=(11, 6))
    i = m["idx"]
    com = (1 + m["r_com"]).cumprod() - 1
    sem = (1 + m["r_sem"]).cumprod() - 1
    ipca = m["ipca_idx"] / m["ipca_idx"].iloc[0] - 1
    ax.plot(i, com * 100, label="Carteira em R$ (com câmbio)", linewidth=2)
    ax.plot(i, sem * 100, label="Carteira sem variação cambial", linewidth=2, linestyle="--")
    ax.plot(i, ipca * 100, label="IPCA acumulado", linewidth=1.6, color="grey")
    ax.axhline(0, color="grey", linewidth=0.8)
    ax.set_ylabel("Rentabilidade acumulada (%)")
    ax.set_title(titulo, fontsize=13, weight="bold")
    ax.grid(alpha=0.3)
    ax.legend(frameon=False)
    import matplotlib.dates as mdates
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d/%m/%y"))
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(arquivo, dpi=150, bbox_inches="tight")
    plt.close(fig)


# =========================================================================================
# 9. Comandos
# =========================================================================================
FONTES = None


def fontes():
    global FONTES
    if FONTES is None:
        FONTES = FontesReais()
    return FONTES


def proximo_id(c):
    return max([m["id"] for m in c["movimentos"]], default=0) + 1


def obter(c, nome):
    t = nome.upper()
    if t not in c["ativos"]:
        sys.exit(f"'{t}' não está na carteira. Veja os ativos com:  python tracker.py carteira")
    return t, c["ativos"][t]


def selecionar(c, args):
    ts = list(c["ativos"])
    if getattr(args, "ativo", None):
        t, _ = obter(c, args.ativo)
        return [t]
    if getattr(args, "classe", None):
        ts = [t for t in ts if c["ativos"][t]["classe"] == args.classe]
    return ts


def _fx_informado(args):
    return float(args.cambio) if getattr(args, "cambio", None) else None


def novo_mov(c, **kw):
    m = {"id": proximo_id(c), **kw}
    c["movimentos"].append(m)
    return m


def buscar_setores_yahoo(a, t):
    f = fontes()
    try:
        if a["classe"] == "acao":
            chave = normalizar_setor(f.info(t)["setor"])
            if chave:
                a.update(setores={chave: 1.0}, fonte_setores="Yahoo Finance (setor da empresa)",
                         data_setores=hoje().isoformat())
                return True
        else:
            pesos = f.pesos_fundo(t)
            if pesos:
                a["setores"], _ = normalizar_pesos(pesos)
                a.update(fonte_setores="Yahoo Finance (funds_data) — sem data de referência",
                         data_setores=hoje().isoformat())
                return True
    except Exception as e:
        print(f"  (não consegui consultar o Yahoo: {e.__class__.__name__})")
    return False


def cmd_adicionar(args, c):
    t = args.ticker.upper()
    if t in c["ativos"]:
        sys.exit(f"'{t}' já existe. Para comprar mais: python tracker.py comprar {t} --quantidade ... --preco ...")
    d = data(args.data) if args.data else hoje()
    classe = args.classe
    if classe in ("fundo", "outro"):
        if args.valor is None:
            sys.exit("Para fundo/outro informe --valor (quanto você aplicou).")
        moeda = (args.moeda or "BRL").upper()
        c["ativos"][t] = {"nome": args.nome or t, "classe": classe, "moeda": moeda, "fonte_preco": "manual",
                          "setores": {}, "fonte_setores": None, "data_setores": None, "retencao_dividendos": 0.0}
        m = novo_mov(c, data=d.isoformat(), ativo=t, tipo="aporte", valor=args.valor,
                     cambio=_fx_informado(args), fonte_cambio="informado" if args.cambio else "PTAX do dia")
        c["valores"].setdefault(t, []).append({"data": d.isoformat(), "valor": args.valor})
    else:
        if args.quantidade is None:
            sys.exit("Informe --quantidade (e --preco pago; sem --preco uso o fechamento do dia).")
        f = fontes()
        info = f.info(t)
        if not classe:
            classe = QUOTE_TYPES.get(info["quote_type"], "acao")
            if classe == "etf" and "." in t and t.rsplit(".", 1)[1] in SUFIXOS_EUROPA:
                classe = "ucits"
                print(f"  {t} é um ETF listado na Europa: cadastrado como UCITS ETF (sem retenção americana).")
        px = f.precos(t, d - dt.timedelta(days=10), max(d, hoje()))
        moeda = px.attrs.get("moeda") or info["moeda"] or "USD"
        if args.moeda and args.moeda.upper() != moeda:
            print(f"  Aviso: {t} é cotado em {moeda} no Yahoo; usando {moeda} (não {args.moeda.upper()}).")
        preco = args.preco
        if preco is None:
            fech = px["Close"][px.index <= pd.Timestamp(d)]
            if fech.empty:
                sys.exit("Sem cotação para essa data; informe --preco.")
            preco = float(fech.iloc[-1])
        ret = args.retencao if args.retencao is not None else (0.30 if moeda == "USD" and classe in ("acao", "etf") else 0.0)
        c["ativos"][t] = {"nome": args.nome or info["nome"], "classe": classe, "moeda": moeda, "fonte_preco": "yahoo",
                          "setores": {}, "fonte_setores": None, "data_setores": None, "retencao_dividendos": ret}
        novo_mov(c, data=d.isoformat(), ativo=t, tipo="compra", quantidade=args.quantidade, preco=preco,
                 corretagem=args.corretagem or 0.0, cambio=_fx_informado(args),
                 fonte_cambio="informado" if args.cambio else "PTAX do dia")
    a = c["ativos"][t]
    if args.setores:
        aplicar_setores_texto(a, args.setores)
    elif a["fonte_preco"] == "yahoo":
        buscar_setores_yahoo(a, t)
    print(f"✓ {t} — {a['nome']}  [{CLASSES[a['classe']]}, {a['moeda']}]"
          + (f"  retenção de dividendos: {a['retencao_dividendos']:.0%}" if a["fonte_preco"] == "yahoo" else ""))
    mostrar_setores(t, a)
    if not a["setores"] and a["classe"] in ("etf", "ucits", "fundo"):
        print("  ⚠ Sem composição setorial. Opções (da mais para a menos confiável):\n"
              f"    python tracker.py setores-arquivo {t} posicoes.csv      (arquivo 'Download Holdings' do emissor)\n"
              f"    python tracker.py setores-ishares {t} LINK_DA_PAGINA   (só iShares)\n"
              f"    python tracker.py setores-lamina {t} lamina.pdf\n"
              f'    python tracker.py definir-setores {t} "Tecnologia=40; Financeiro=20"')


def cmd_adicionar_rf(args, c):
    t = args.nome.upper()
    if t in c["ativos"]:
        sys.exit(f"'{t}' já existe. Para aplicar mais: python tracker.py aportar {t} --valor ...")
    ind = args.indexador
    moeda = (args.moeda or "BRL").upper()
    if ind != "pre" and moeda != "BRL":
        sys.exit("Só títulos prefixados podem ser em moeda estrangeira (CDI, Selic e IPCA são brasileiros).")
    rf = {"indexador": ind, "vencimento": args.vencimento, "isento": args.isento, "emissor": args.emissor}
    if ind == "pre":
        if args.taxa is None:
            sys.exit("Prefixado: informe --taxa (ex.: --taxa 13.5 para 13,5% a.a.).")
        rf.update(taxa=args.taxa / 100, base=args.base or (252 if moeda == "BRL" else 365))
    elif ind == "ipca":
        if args.taxa is None:
            sys.exit("IPCA+: informe --taxa real (ex.: --taxa 6.8 para IPCA + 6,8% a.a.).")
        rf.update(taxa=args.taxa / 100)
    else:
        rf.update(percentual=(args.percentual or 100) / 100, spread=(args.spread or 0) / 100)
    c["ativos"][t] = {"nome": args.descricao or t, "classe": "renda_fixa", "moeda": moeda, "fonte_preco": "renda_fixa",
                      "setores": {"renda_fixa": 1.0}, "fonte_setores": "Renda fixa", "data_setores": None,
                      "retencao_dividendos": 0.0, "rf": rf}
    d = data(args.data) if args.data else hoje()
    novo_mov(c, data=d.isoformat(), ativo=t, tipo="aporte", valor=args.valor, cambio=_fx_informado(args),
             fonte_cambio="informado" if args.cambio else "PTAX do dia")
    desc = {"pre": lambda: f"{br(args.taxa)}% a.a.", "ipca": lambda: f"IPCA + {br(args.taxa)}% a.a.",
            "cdi": lambda: f"{br(args.percentual or 100, 0)}% do CDI" + (f" + {br(args.spread)}% a.a." if args.spread else ""),
            "selic": lambda: f"{br(args.percentual or 100, 0)}% da Selic" + (f" + {br(args.spread)}% a.a." if args.spread else "")}
    print(f"✓ {t}: {INDEXADORES[ind]} — {desc[ind]()} — aplicado {br(args.valor)} {moeda} em {d:%d/%m/%Y}"
          + (" (isento de IR)" if args.isento else ""))


def cmd_comprar_vender(args, c, tipo):
    t, a = obter(c, args.ticker)
    if a["fonte_preco"] != "yahoo":
        sys.exit(f"{t} não é ativo de bolsa; use 'aportar' ou 'resgatar'.")
    d = data(args.data) if args.data else hoje()
    preco = args.preco
    if preco is None:
        px = fontes().precos(t, d - dt.timedelta(days=10), max(d, hoje()))
        preco = float(px["Close"][px.index <= pd.Timestamp(d)].iloc[-1])
    m = novo_mov(c, data=d.isoformat(), ativo=t, tipo=tipo, quantidade=args.quantidade, preco=preco,
                 corretagem=args.corretagem or 0.0, cambio=_fx_informado(args),
                 fonte_cambio="informado" if args.cambio else "PTAX do dia")
    nomes = {"compra": "Compra", "venda": "Venda", "reinvestir": "Reinvestimento de dividendos"}
    print(f"✓ {nomes[tipo]} #{m['id']}: {br(args.quantidade, 4)} × {br(preco)} {a['moeda']} de {t} em {d:%d/%m/%Y}")


def cmd_aportar_resgatar(args, c, tipo):
    t, a = obter(c, args.nome)
    if a["fonte_preco"] == "yahoo":
        sys.exit(f"{t} é ativo de bolsa; use 'comprar' ou 'vender'.")
    d = data(args.data) if args.data else hoje()
    m = novo_mov(c, data=d.isoformat(), ativo=t, tipo=tipo, valor=args.valor, cambio=_fx_informado(args),
                 fonte_cambio="informado" if args.cambio else "PTAX do dia")
    print(f"✓ {'Aporte' if tipo == 'aporte' else 'Resgate'} #{m['id']} em {t}: {br(args.valor)} {a['moeda']} em {d:%d/%m/%Y}")
    if a["fonte_preco"] == "manual":
        print(f"  Lembre de informar o valor total atualizado:  python tracker.py valor {t} --valor ...")


def cmd_valor(args, c):
    t, a = obter(c, args.nome)
    if a["fonte_preco"] != "manual":
        sys.exit(f"{t} é calculado automaticamente ({a['fonte_preco']}); não precisa informar valor.")
    d = (data(args.data) if args.data else hoje()).isoformat()
    lista = [v for v in c["valores"].setdefault(t, []) if v["data"] != d]
    lista.append({"data": d, "valor": args.valor})
    c["valores"][t] = sorted(lista, key=lambda v: v["data"])
    print(f"✓ {t}: valor de {br(args.valor)} {a['moeda']} em {data(d):%d/%m/%Y}")


def cmd_movimentos(args, c):
    ms = sorted(c["movimentos"], key=lambda m: (m["data"], m["id"]))
    if args.ativo:
        t, _ = obter(c, args.ativo)
        ms = [m for m in ms if m["ativo"] == t]
    linhas = []
    for m in ms:
        moeda = c["ativos"].get(m["ativo"], {}).get("moeda", "")
        qp = f"{br(m['quantidade'], 4)} × {br(m['preco'])}" if m.get("quantidade") else br(m.get("valor", 0))
        linhas.append([m["id"], data(m["data"]).strftime("%d/%m/%Y"), m["ativo"], m["tipo"], qp + f" {moeda}",
                       br(m.get("corretagem", 0) or 0), br(m["cambio"], 4) if m.get("cambio") else "PTAX"])
    tabela(["#", "Data", "Ativo", "Tipo", "Qtd × Preço / Valor", "Corretagem", "Câmbio"], linhas, direita={0, 4, 5, 6})


def cmd_editar_movimento(args, c):
    m = next((m for m in c["movimentos"] if m["id"] == args.id), None)
    if not m:
        sys.exit(f"Movimento #{args.id} não existe. Veja com: python tracker.py movimentos")
    for campo in ("quantidade", "preco", "corretagem", "valor", "cambio"):
        v = getattr(args, campo)
        if v is not None:
            m[campo] = v
            if campo == "cambio":
                m["fonte_cambio"] = "informado"
    if args.data:
        m["data"] = data(args.data).isoformat()
    print(f"✓ Movimento #{m['id']} atualizado.")


def cmd_remover_movimento(args, c):
    antes = len(c["movimentos"])
    c["movimentos"] = [m for m in c["movimentos"] if m["id"] != args.id]
    print("✓ Removido." if len(c["movimentos"]) < antes else f"Movimento #{args.id} não existe.")


def cmd_remover(args, c):
    t, _ = obter(c, args.nome)
    del c["ativos"][t]
    c["movimentos"] = [m for m in c["movimentos"] if m["ativo"] != t]
    c["valores"].pop(t, None)
    print(f"✓ {t} e seus movimentos removidos.")


# ---- setores ---------------------------------------------------------------------------
def mostrar_setores(t, a):
    if a.get("setores"):
        itens = sorted(a["setores"].items(), key=lambda x: -x[1])
        print("  Setores: " + ", ".join(f"{SETORES_PT[k]} {br(v * 100, 1)}%" for k, v in itens[:6])
              + (" …" if len(itens) > 6 else ""))
        print(f"  Fonte: {a.get('fonte_setores')}" + (f" — obtido em {data(a['data_setores']):%d/%m/%Y}"
                                                       if a.get("data_setores") else ""))


def aplicar_setores_texto(a, texto, fonte="Informado manualmente"):
    pares = re.findall(r"([^=;]+?)\s*=\s*(-?[\d.,]+)", texto)
    if not pares:
        sys.exit('Formato: "Tecnologia=47.6; Financeiro=15.4; Energia=6.3"')
    brutos = {n.strip(" ,;"): float(v.strip(",.").replace(",", ".")) for n, v in pares}
    a["setores"], desc = normalizar_pesos(brutos)
    a.update(fonte_setores=fonte, data_setores=hoje().isoformat())
    if desc:
        print(f"  Aviso: não reconheci {', '.join(desc)} — contados em 'Outros'.")


def _aplicar_pesos(t, a, pesos, fonte, data_ref, n):
    a["setores"], desc = normalizar_pesos(pesos)
    a.update(fonte_setores=fonte + (f" (posições de {data_ref})" if data_ref else ""), data_setores=hoje().isoformat())
    print(f"✓ {t}: {n} posições lidas.")
    mostrar_setores(t, a)
    if desc:
        print(f"  Setores não reconhecidos (em 'Outros'): {', '.join(sorted(set(desc)))}")


def cmd_setores_arquivo(args, c):
    t, a = obter(c, args.ticker)
    p = Path(args.arquivo)
    pesos, data_ref, n = ler_posicoes(p.read_bytes(), p.name)
    _aplicar_pesos(t, a, pesos, f"Arquivo de posições {p.name}", data_ref, n)


def cmd_setores_ishares(args, c):
    t, a = obter(c, args.ticker)
    url = url_ishares_csv(args.url, t)
    print(f"  Baixando {url}")
    try:
        bruto = fontes().baixar(url)
    except Exception as e:
        sys.exit(f"Não consegui baixar ({e}). Abra a página, clique em 'Download Holdings' (CSV) e use:\n"
                 f"  python tracker.py setores-arquivo {t} arquivo.csv")
    if bruto[:200].lstrip().lower().startswith((b"<!doctype", b"<html")):
        sys.exit("A iShares devolveu uma página em vez do CSV (bloqueio ou link errado). Baixe manualmente\n"
                 f"em 'Download Holdings' e use:  python tracker.py setores-arquivo {t} arquivo.csv")
    pesos, data_ref, n = ler_posicoes(bruto, "ishares.csv")
    _aplicar_pesos(t, a, pesos, "iShares — arquivo oficial de posições", data_ref, n)


def cmd_setores_lamina(args, c):
    t, a = obter(c, args.ticker)
    try:
        import pdfplumber
    except ImportError:
        sys.exit("Instale:  pip install pdfplumber")
    with pdfplumber.open(args.pdf) as pdf:
        texto = "\n".join(p.extract_text() or "" for p in pdf.pages)
    achados = setores_de_texto_pdf(texto)
    if len(achados) < 3:
        sys.exit("Esta lâmina não tem a tabela de setores em texto (em muitas, como as da iShares, os setores\n"
                 "são um gráfico). Use o arquivo de posições do emissor:\n"
                 f"  python tracker.py setores-arquivo {t} posicoes.csv")
    a["setores"], _ = normalizar_pesos(achados)
    a.update(fonte_setores=f"Lâmina {Path(args.pdf).name}", data_setores=hoje().isoformat())
    mostrar_setores(t, a)
    print("  Confira com a lâmina; corrija com 'definir-setores' se preciso.")


def cmd_definir_setores(args, c):
    t, a = obter(c, args.ticker)
    aplicar_setores_texto(a, args.setores)
    mostrar_setores(t, a)


def cmd_setores_info(args, c):
    linhas = []
    for t, a in c["ativos"].items():
        idade = (hoje() - data(a["data_setores"])).days if a.get("data_setores") else None
        alerta = "⚠ atualizar" if (idade is not None and idade > 90 and a["classe"] in ("etf", "ucits", "fundo")) \
            else ("⚠ sem dados" if not a.get("setores") else "")
        linhas.append([t, CLASSES[a["classe"]], a.get("fonte_setores") or "—",
                       f"{idade} dias" if idade is not None else "—", alerta])
    tabela(["Ativo", "Classe", "Fonte dos setores", "Idade", ""], linhas, direita={3})


def cmd_ipca_projecao(args, c):
    m = re.fullmatch(r"(\d{4})-(\d{2})", args.mes)
    if not m:
        sys.exit("Use o formato AAAA-MM, ex.: 2026-09")
    proj = c["config"].setdefault("ipca_projecoes", {})
    if args.taxa is None:
        proj.pop(args.mes, None)
        print(f"✓ Projeção de {args.mes} removida (volta a usar IBGE/Focus).")
    else:
        proj[args.mes] = args.taxa
        print(f"✓ IPCA de {args.mes} projetado em {br(args.taxa)}% (vale até sair o dado oficial do IBGE).")


# ---- relatórios ------------------------------------------------------------------------
def _motor(c, args=None):
    ate = data(args.ate) if args is not None and getattr(args, "ate", None) else None
    return Motor(c, fontes(), ate)


def _avisos(motor):
    for a in dict.fromkeys(fontes().avisos + motor.avisos):
        print(f"  ⚠ {a}")


def cmd_carteira(args, c):
    motor = _motor(c, args)
    ts = selecionar(c, args)
    linhas, total = [], 0.0
    info = []
    for t in ts:
        df = motor.serie(t)
        if df is None:
            continue
        a = c["ativos"][t]
        ult = df.iloc[-1]
        v_brl = float(ult["v_brl"])
        total += v_brl
        info.append((t, a, ult, v_brl))
    for t, a, ult, v_brl in sorted(info, key=lambda x: -x[3]):
        extra = ""
        if a["fonte_preco"] == "yahoo":
            extra = f"{br(ult['qtd'], 4)} × {br(ult['preco'])}"
            if ult["caixa_div"] > 0.005:
                extra += f" + div. {br(ult['caixa_div'])}"
        elif a["fonte_preco"] == "renda_fixa":
            ir = motor.ir_estimado_rf(t)
            extra = f"IR estimado {br(ir)}" if ir else "isento/sem IR"
        else:
            ult_rel = max([v["data"] for v in c["valores"].get(t, [])], default=None)
            if ult_rel:
                dias = (motor.ate - data(ult_rel)).days
                extra = f"valor de {data(ult_rel):%d/%m}" + (" ⚠ desatualizado" if dias > 10 else "")
        linhas.append([t, CLASSES[a["classe"]], f"{br(ult['v_nat'])} {a['moeda']}", br(ult["fx"], 4) if a["moeda"] != "BRL" else "—",
                       rs(v_brl), pct(v_brl / total if total else 0, sinal=False), extra])
    print(f"\nCarteira em {motor.ate:%d/%m/%Y}\n")
    tabela(["Ativo", "Classe", "Valor (moeda)", "Câmbio", "Valor R$", "%", "Detalhe"], linhas, direita={2, 3, 4, 5})
    print(f"\nTotal: {rs(total)}")
    _avisos(motor)


def _imprimir_metricas(m, titulo):
    print(f"\n{titulo}\nPeríodo: {m['inicio']:%d/%m/%Y} → {m['fim']:%d/%m/%Y} ({m['dias']} dias)\n")
    nota_anual = "" if m["anual_com"] is not None else \
        "  (só com 1 ano+, padrão GIPS; use --anualizar-curto)"
    nota_vol = f"  ({m['semanas']} semanas)" if m["vol_com"] is not None else f"  (precisa de 4+ semanas; há {m['semanas']})"
    linhas = [
        ["Rentabilidade no período (TWR)", pct(m["twr_com"]), pct(m["twr_sem"]), ""],
        [f"Descontada a inflação (IPCA {pct(m['ipca'], False)})", pct(m["real_com"]), pct(m["real_sem"]), ""],
        ["Anualizada", pct(m["anual_com"]), pct(m["anual_sem"]), nota_anual],
        ["Anualizada, descontada a inflação", pct(m["anual_real_com"]), pct(m["anual_real_sem"]), ""],
        ["Volatilidade anualizada", pct(m["vol_com"], False), pct(m["vol_sem"], False), nota_vol],
    ]
    tabela(["", "Com câmbio", "Sem câmbio", ""], linhas, direita={1, 2})
    print()
    tabela(["Em dinheiro", "Valor", ""], [
        ["Valor atual", rs(m["valor_brl"]), ""],
        ["Aportes no período", rs(m["aportes"]), ""],
        ["Resultado (com câmbio)", rs(m["resultado"]), ""],
        ["Resultado sem variação cambial", rs(m["resultado_sem_cambio"]), "câmbio travado no que você pagou em cada ativo"],
        ["Efeito do câmbio", rs(m["efeito_cambio"]), ""],
        ["Resultado real (acima do IPCA)", rs(m["resultado_real"]), "cada aporte corrigido pelo IPCA desde a sua data"],
    ], direita={1})


def _periodo(args, motor):
    if args.desde:
        return data(args.desde)
    fim = motor.ate
    return {"mes": dt.date(fim.year, fim.month, 1) - dt.timedelta(days=1),
            "ano": dt.date(fim.year - 1, 12, 31), "12m": fim - dt.timedelta(days=365)}.get(args.periodo)


def cmd_rentabilidade(args, c):
    motor = _motor(c, args)
    ts = selecionar(c, args)
    desde = _periodo(args, motor)
    amin = 30 if args.anualizar_curto else 365
    m = calcular(motor, ts, desde, amin)
    if m is None:
        sys.exit("Nenhum ativo com movimentos nesse filtro.")
    alvo = args.ativo.upper() if args.ativo else (CLASSES_PLURAL[args.classe] if args.classe else "Carteira inteira")
    _imprimir_metricas(m, f"Rentabilidade — {alvo}")
    if len(ts) > 1:
        print("\nPor ativo:\n")
        linhas = []
        for t in ts:
            mi = calcular(motor, [t], desde, amin)
            if mi is None:
                continue
            linhas.append([t, pct(mi["twr_com"]), pct(mi["twr_sem"]), pct(mi["real_com"]), pct(mi["anual_real_com"]),
                           pct(mi["vol_com"], False), rs(mi["resultado"]), rs(mi["efeito_cambio"])])
        tabela(["Ativo", "Rent. R$", "Sem câmbio", "Real", "Real anual", "Volat.", "Resultado", "Efeito câmbio"], linhas)
        if amin == 30:
            print("\n  ⚠ Anualização de períodos curtos projeta o resultado para 1 ano; use só para comparar com taxas.")
    meses_proj = sorted(k for k, v in motor.fonte_ipca.items() if v != "IBGE" and k <= motor.ate.strftime("%Y-%m"))
    if meses_proj:
        print(f"\n  IPCA estimado para {', '.join(meses_proj)} ({', '.join(sorted({motor.fonte_ipca[k] for k in meses_proj}))}); "
              "o número muda quando o IBGE divulgar.")
    _avisos(motor)


def cmd_alocacao(args, c):
    motor = _motor(c, args)
    ts = selecionar(c, args)
    for por in (["classe", "setor"] if args.por == "ambos" else [args.por]):
        vals = alocacao(motor, ts, por)
        total = sum(vals.values())
        print(f"\nAlocação por {'classe de ativo' if por == 'classe' else 'setor (ETFs abertos pela composição)'}\n")
        tabela([por.capitalize(), "Valor R$", "%", ""],
               [[k, rs(v), pct(v / total, False), "█" * round(v / total * 30)] for k, v in vals.items()], direita={1, 2})
        print(f"Total: {rs(total)}")
    sem = [t for t in ts if not c["ativos"][t].get("setores")]
    if sem and args.por != "classe":
        print(f"\n  ⚠ Sem composição setorial (em 'Outros'): {', '.join(sem)}")
    _avisos(motor)


def cmd_grafico(args, c):
    motor = _motor(c, args)
    ts = selecionar(c, args)
    PASTA_GRAFICOS.mkdir(parents=True, exist_ok=True)
    sufixo = f"_{args.classe}" if args.classe else ""
    arq = Path(args.arquivo) if args.arquivo else PASTA_GRAFICOS / f"{args.tipo}{sufixo}_{motor.ate:%Y%m%d}.png"
    rot = f" — {CLASSES[args.classe]}" if args.classe else ""
    if args.tipo == "pizza-classes":
        grafico_pizza(alocacao(motor, ts, "classe"), f"Carteira por classe de ativo{rot} — {motor.ate:%d/%m/%Y}", arq)
    elif args.tipo == "pizza-setores":
        grafico_pizza(alocacao(motor, ts, "setor"), f"Carteira por setor{rot} — {motor.ate:%d/%m/%Y}", arq)
    elif args.tipo == "setores-mensal":
        df = retorno_mensal_setores(motor, ts, cambio=not args.sem_cambio)
        if df.empty:
            sys.exit("Ainda não há dados suficientes.")
        if not args.todos_setores:
            aloc = alocacao(motor, ts, "setor")
            tot = sum(aloc.values())
            manter = [s_ for s_ in df.columns if s_ == "Carteira" or aloc.get(s_, 0) / tot >= 0.03]
            fora = [s_ for s_ in df.columns if s_ not in manter]
            df = df[manter]
            if fora:
                print(f"  (setores com menos de 3% da carteira fora do gráfico: {', '.join(fora)}; use --todos-setores)")
        titulo = ("Rentabilidade acumulada" if args.acumulado else "Rentabilidade mensal") + " por setor" + \
                 (" (sem câmbio)" if args.sem_cambio else " (em R$)") + rot
        grafico_linhas(df, titulo, arq, "Rentabilidade acumulada" if args.acumulado else "Rentabilidade no mês",
                       acumulado=args.acumulado,
                       nota="Setores de ETFs usam o retorno do ETF inteiro (aproximação); ações entram 100% no setor.")
        print("\nRentabilidade mensal por setor:\n")
        tabela(["Mês"] + list(df.columns), [[pd.Period(i).strftime("%m/%Y")] + [pct(v) for v in row]
                                            for i, row in df.iterrows()])
    elif args.tipo == "evolucao":
        m = calcular(motor, ts)
        grafico_evolucao(m, f"Rentabilidade acumulada{rot} × IPCA", arq)
    print(f"\n✓ Gráfico salvo em {arq.resolve()}")
    _avisos(motor)


def cmd_atualizar(args, c):
    f = fontes()
    f.forcar = True
    motor = Motor(c, f)
    print(f"Atualizando dados de {motor.ate:%d/%m/%Y}...\n")
    ok = 0
    for t, a in c["ativos"].items():
        try:
            df = motor.serie(t)
            if df is None:
                continue
            ult = df.iloc[-1]
            ok += 1
            print(f"✓ {t:<12} {br(ult['v_nat']):>14} {a['moeda']}  → {rs(float(ult['v_brl'])):>16}")
        except Exception as e:
            print(f"✗ {t:<12} erro: {e}")
    usd = motor._fx.get("USD")
    if usd is not None and len(usd):
        print(f"\nPTAX USD mais recente: {br(float(usd.iloc[-1]), 4)} em {usd.index[-1]:%d/%m/%Y}")
    if motor.fonte_ipca:
        ult_ibge = max((k for k, v in motor.fonte_ipca.items() if v == "IBGE"), default=None)
        if ult_ibge:
            print(f"Último IPCA oficial: {ult_ibge}")
    print("\nPendências:")
    pend = 0
    for t, a in c["ativos"].items():
        if a["fonte_preco"] == "manual":
            ult_rel = max([v["data"] for v in c["valores"].get(t, [])], default=None)
            if not ult_rel or (motor.ate - data(ult_rel)).days > 7:
                print(f"  • {t}: informe o valor atual →  python tracker.py valor {t} --valor ...")
                pend += 1
        if a["classe"] in ("etf", "ucits", "fundo"):
            idade = (motor.ate - data(a["data_setores"])).days if a.get("data_setores") else None
            if idade is None or idade > 90:
                print(f"  • {t}: composição setorial {'ausente' if idade is None else f'com {idade} dias'} → setores-arquivo/ishares")
                pend += 1
    if not pend:
        print("  nenhuma")
    _avisos(motor)


# ---- exportação Excel -------------------------------------------------------------------
def cmd_exportar(args, c):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    motor = _motor(c, args)
    ts = [t for t in c["ativos"] if motor.serie(t) is not None]
    F = "Arial"
    H, HF = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    N, NOTA = Font(name=F), Font(name=F, italic=True, size=9, color="808080")
    MO, PC = '#,##0.00;(#,##0.00);"-"', '0.00%;(0.00%);"-"'
    wb = Workbook()

    def aba(nome, cab, larg, primeira=False):
        ws = wb.active if primeira else wb.create_sheet(nome)
        ws.title = nome
        for i, (h, w) in enumerate(zip(cab, larg), start=1):
            x = ws.cell(row=1, column=i, value=h)
            x.font, x.fill = H, HF
            x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "B2"
        return ws

    ws = aba("Posições", ["Ativo", "Nome", "Classe", "Moeda", "Valor (moeda)", "Câmbio", "Valor R$", "% carteira",
                          "Aportes R$", "Resultado R$", "Rent. R$ (TWR)", "Rent. sem câmbio", "Rent. real",
                          "Real anualizada", "Volatilidade", "Efeito câmbio R$"],
             [10, 30, 12, 8, 14, 9, 14, 10, 14, 14, 12, 12, 11, 12, 11, 14], primeira=True)
    n = len(ts)
    for i, t in enumerate(ts, start=2):
        a, mi = c["ativos"][t], calcular(motor, [t])
        ult = motor.serie(t).iloc[-1]
        vals = [t, a["nome"], CLASSES[a["classe"]], a["moeda"], float(ult["v_nat"]), float(ult["fx"]), None, None,
                mi["aportes"], None, mi["twr_com"], mi["twr_sem"], mi["real_com"], mi["anual_real_com"],
                mi["vol_com"], mi["efeito_cambio"]]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v).font = N
        ws.cell(row=i, column=7, value=f"=E{i}*F{i}")
        ws.cell(row=i, column=8, value=f'=IF(SUM($G$2:$G${n + 1})=0,"",G{i}/SUM($G$2:$G${n + 1}))')
        ws.cell(row=i, column=10, value=f"=G{i}-I{i}")
        for j, fm in ((5, MO), (6, "0.0000"), (7, MO), (8, PC), (9, MO), (10, MO), (11, PC), (12, PC), (13, PC),
                      (14, PC), (15, PC), (16, MO)):
            ws.cell(row=i, column=j).number_format = fm
    tr = n + 2
    ws.cell(row=tr, column=1, value="TOTAL").font = Font(name=F, bold=True)
    for j in (7, 9, 10, 16):
        L = get_column_letter(j)
        x = ws.cell(row=tr, column=j, value=f"=SUM({L}2:{L}{n + 1})")
        x.number_format, x.font = MO, Font(name=F, bold=True)
    ws.cell(row=tr + 2, column=1, value=(f"Rentabilidades calculadas pelo tracker.py em {motor.ate:%d/%m/%Y} a partir da "
                                        "série diária (TWR). Resultado = valor atual − aportes líquidos em R$ pagos. "
                                        "Câmbio: PTAX/BCB; IPCA: IBGE/BCB + Focus.")).font = NOTA

    for nome, por in (("Alocação Classes", "classe"), ("Alocação Setores", "setor")):
        vals = alocacao(motor, ts, por)
        w = aba(nome, [por.capitalize(), "Valor R$", "%"], [30, 16, 10])
        k = len(vals)
        for i, (rot, v) in enumerate(vals.items(), start=2):
            w.cell(row=i, column=1, value=rot).font = N
            w.cell(row=i, column=2, value=v).number_format = MO
            w.cell(row=i, column=3, value=f"=B{i}/SUM($B$2:$B${k + 1})").number_format = PC

    df = retorno_mensal_setores(motor, ts)
    w = aba("Setores Mensal", ["Mês"] + list(df.columns), [10] + [14] * len(df.columns))
    for i, (mes, row) in enumerate(df.iterrows(), start=2):
        w.cell(row=i, column=1, value=str(mes))
        for j, v in enumerate(row, start=2):
            x = w.cell(row=i, column=j, value=None if pd.isna(v) else float(v))
            x.number_format = PC

    m = calcular(motor, ts)
    w = aba("Série Diária", ["Data", "Valor R$", "Fluxo R$", "Índice TWR (com câmbio)", "Índice TWR (sem câmbio)",
                             "Índice IPCA"], [12, 14, 14, 16, 16, 12])
    niv_c, niv_s = (1 + m["r_com"]).cumprod(), (1 + m["r_sem"]).cumprod()
    for i, d in enumerate(m["idx"], start=2):
        w.cell(row=i, column=1, value=d.date()).number_format = "DD/MM/YYYY"
        w.cell(row=i, column=2, value=float(m["V"].loc[d])).number_format = MO
        w.cell(row=i, column=3, value=float(m["CF"].loc[d])).number_format = MO
        w.cell(row=i, column=4, value=float(niv_c.loc[d])).number_format = "0.0000"
        w.cell(row=i, column=5, value=float(niv_s.loc[d])).number_format = "0.0000"
        w.cell(row=i, column=6, value=float(m["ipca_idx"].loc[d] / m["ipca_idx"].iloc[0])).number_format = "0.0000"

    w = aba("Movimentos", ["#", "Data", "Ativo", "Tipo", "Quantidade", "Preço", "Valor", "Corretagem", "Câmbio"],
            [6, 12, 10, 12, 12, 12, 12, 12, 10])
    for i, mv in enumerate(sorted(c["movimentos"], key=lambda x: (x["data"], x["id"])), start=2):
        for j, v in enumerate([mv["id"], data(mv["data"]), mv["ativo"], mv["tipo"], mv.get("quantidade"),
                               mv.get("preco"), mv.get("valor"), mv.get("corretagem"), mv.get("cambio")], start=1):
            x = w.cell(row=i, column=j, value=v)
            if j == 2:
                x.number_format = "DD/MM/YYYY"
    for wsx in wb.worksheets:
        for row in wsx.iter_rows():
            for x in row:
                if x.font.name != F:
                    x.font = Font(name=F, bold=x.font.bold, italic=x.font.italic, color=x.font.color, size=x.font.size)
    destino = Path(args.arquivo)
    wb.save(destino)
    print(f"✓ Relatório salvo em {destino.resolve()}")


# =========================================================================================
# 10. Linha de comando
# =========================================================================================
def construir_parser():
    p = argparse.ArgumentParser(description="Rastreador de carteira (v2). Guia completo no README.md.")
    p.add_argument("--carteira", default=str(ARQUIVO_PADRAO), help="arquivo JSON da carteira")
    sub = p.add_subparsers(dest="cmd", required=True)
    classes_mercado = ["acao", "etf", "ucits", "fundo", "outro"]

    s = sub.add_parser("adicionar", help="cadastra ação, ETF, UCITS ETF, fundo ou outro ativo")
    s.add_argument("ticker", help="ticker do Yahoo (VLO, VOO, EMVL.L) ou um nome curto (fundos)")
    s.add_argument("--classe", choices=classes_mercado)
    s.add_argument("--quantidade", type=float)
    s.add_argument("--preco", type=float, help="preço pago por unidade (na moeda do ativo)")
    s.add_argument("--corretagem", type=float, default=0.0)
    s.add_argument("--valor", type=float, help="valor aplicado (fundo/outro)")
    s.add_argument("--moeda", help="USD ou BRL (fundo/outro; ações usam a moeda do Yahoo)")
    s.add_argument("--cambio", type=float, help="R$ por unidade da moeda que você pagou (padrão: PTAX do dia)")
    s.add_argument("--retencao", type=float, help="imposto retido nos dividendos (padrão 0.30 para EUA)")
    s.add_argument("--data", help="AAAA-MM-DD ou DD/MM/AAAA (padrão: hoje)")
    s.add_argument("--nome")
    s.add_argument("--setores", help='ex.: "Tecnologia=40; Saúde=20"')

    s = sub.add_parser("adicionar-rf", help="cadastra renda fixa (prefixado, CDI, Selic, IPCA+)")
    s.add_argument("nome", help="nome curto, ex.: CDB_BANCO_X, TESOURO_IPCA_2035")
    s.add_argument("--indexador", required=True, choices=list(INDEXADORES))
    s.add_argument("--valor", type=float, required=True)
    s.add_argument("--data", help="data da aplicação")
    s.add_argument("--taxa", type=float, help="%% a.a. (prefixado: taxa total; IPCA+: taxa real)")
    s.add_argument("--percentual", type=float, help="%% do CDI/Selic (padrão 100)")
    s.add_argument("--spread", type=float, help="%% a.a. somado ao CDI/Selic (ex.: CDI + 1)")
    s.add_argument("--vencimento")
    s.add_argument("--isento", action="store_true", help="LCI, LCA, CRI, CRA, debênture incentivada")
    s.add_argument("--moeda", help="BRL (padrão); USD só para prefixado")
    s.add_argument("--base", type=int, choices=[252, 365], help="contagem de dias do prefixado")
    s.add_argument("--cambio", type=float)
    s.add_argument("--emissor")
    s.add_argument("--descricao")

    for nome, ajuda in (("comprar", "compra de ativo de bolsa"), ("vender", "venda de ativo de bolsa"),
                        ("reinvestir", "compra usando dividendos já recebidos (não é dinheiro novo)")):
        s = sub.add_parser(nome, help=ajuda)
        s.add_argument("ticker")
        s.add_argument("--quantidade", type=float, required=True)
        s.add_argument("--preco", type=float)
        s.add_argument("--corretagem", type=float, default=0.0)
        s.add_argument("--cambio", type=float)
        s.add_argument("--data")

    for nome, ajuda in (("aportar", "novo aporte em fundo/renda fixa/outro"), ("resgatar", "resgate de fundo/renda fixa/outro")):
        s = sub.add_parser(nome, help=ajuda)
        s.add_argument("nome")
        s.add_argument("--valor", type=float, required=True)
        s.add_argument("--cambio", type=float)
        s.add_argument("--data")

    s = sub.add_parser("valor", help="informa o valor atual de um fundo/ativo manual")
    s.add_argument("nome")
    s.add_argument("--valor", type=float, required=True)
    s.add_argument("--data")

    s = sub.add_parser("movimentos", help="lista compras, vendas e aportes")
    s.add_argument("--ativo")
    s = sub.add_parser("editar-movimento", help="corrige um movimento (ex.: câmbio pago)")
    s.add_argument("id", type=int)
    for campo in ("quantidade", "preco", "corretagem", "valor", "cambio"):
        s.add_argument(f"--{campo}", type=float)
    s.add_argument("--data")
    s = sub.add_parser("remover-movimento")
    s.add_argument("id", type=int)
    s = sub.add_parser("remover", help="remove um ativo e seus movimentos")
    s.add_argument("nome")

    s = sub.add_parser("setores-arquivo", help="composição a partir do arquivo de posições do emissor (CSV/XLSX)")
    s.add_argument("ticker")
    s.add_argument("arquivo")
    s = sub.add_parser("setores-ishares", help="baixa as posições direto da página iShares do ETF")
    s.add_argument("ticker")
    s.add_argument("url")
    s = sub.add_parser("setores-lamina", help="lê setores da lâmina PDF (se a tabela estiver em texto)")
    s.add_argument("ticker")
    s.add_argument("pdf")
    s = sub.add_parser("definir-setores", help="informa a composição manualmente")
    s.add_argument("ticker")
    s.add_argument("setores")
    sub.add_parser("setores-info", help="fonte e idade da composição de cada ativo")

    s = sub.add_parser("ipca-projecao", help="sua projeção para um mês ainda sem IPCA oficial")
    s.add_argument("mes", help="AAAA-MM")
    s.add_argument("taxa", type=float, nargs="?", help="%% no mês (omita para apagar)")

    def filtros(s, periodo=False):
        s.add_argument("--classe", choices=list(CLASSES))
        s.add_argument("--ativo")
        s.add_argument("--ate", help="data de corte (padrão: hoje)")
        if periodo:
            s.add_argument("--desde", help="início do período (AAAA-MM-DD)")
            s.add_argument("--periodo", choices=["tudo", "mes", "ano", "12m"], default="tudo")
            s.add_argument("--anualizar-curto", action="store_true",
                           help="anualiza também períodos de 30 dias a 1 ano (fora do padrão GIPS)")

    sub.add_parser("atualizar", help="baixa cotações, câmbio e índices e mostra pendências")
    filtros(sub.add_parser("carteira", help="posição atual de cada ativo"))
    s = sub.add_parser("alocacao", help="%% por classe e por setor")
    filtros(s)
    s.add_argument("--por", choices=["classe", "setor", "ambos"], default="ambos")
    filtros(sub.add_parser("rentabilidade", help="rentabilidade, real, anualizada, volatilidade, câmbio"), periodo=True)
    s = sub.add_parser("grafico", help="pizza-classes | pizza-setores | setores-mensal | evolucao")
    s.add_argument("tipo", choices=["pizza-classes", "pizza-setores", "setores-mensal", "evolucao"])
    filtros(s)
    s.add_argument("--sem-cambio", action="store_true")
    s.add_argument("--acumulado", action="store_true", help="linhas acumuladas em vez de mês a mês")
    s.add_argument("--todos-setores", action="store_true", help="inclui setores com menos de 3%% da carteira")
    s.add_argument("--arquivo")
    s = sub.add_parser("exportar", help="relatório Excel")
    s.add_argument("--arquivo", default="relatorio_carteira.xlsx")
    s.add_argument("--ate")
    return p


def main(argv=None):
    args = construir_parser().parse_args(argv)
    caminho = Path(args.carteira)
    c = carregar(caminho)
    acoes = {
        "adicionar": cmd_adicionar, "adicionar-rf": cmd_adicionar_rf,
        "comprar": lambda a, c: cmd_comprar_vender(a, c, "compra"),
        "vender": lambda a, c: cmd_comprar_vender(a, c, "venda"),
        "reinvestir": lambda a, c: cmd_comprar_vender(a, c, "reinvestir"),
        "aportar": lambda a, c: cmd_aportar_resgatar(a, c, "aporte"),
        "resgatar": lambda a, c: cmd_aportar_resgatar(a, c, "resgate"),
        "valor": cmd_valor, "movimentos": cmd_movimentos, "editar-movimento": cmd_editar_movimento,
        "remover-movimento": cmd_remover_movimento, "remover": cmd_remover,
        "setores-arquivo": cmd_setores_arquivo, "setores-ishares": cmd_setores_ishares,
        "setores-lamina": cmd_setores_lamina, "definir-setores": cmd_definir_setores,
        "setores-info": cmd_setores_info, "ipca-projecao": cmd_ipca_projecao,
        "atualizar": cmd_atualizar, "carteira": cmd_carteira, "alocacao": cmd_alocacao,
        "rentabilidade": cmd_rentabilidade, "grafico": cmd_grafico, "exportar": cmd_exportar,
    }
    acoes[args.cmd](args, c)
    somente_leitura = {"movimentos", "setores-info", "carteira", "alocacao", "rentabilidade", "grafico", "exportar"}
    if args.cmd not in somente_leitura:
        for a in c["ativos"].values():
            a.pop("_lotes", None)
        for m in c["movimentos"]:
            m.pop("_fluxo_nat", None)
        salvar(c, caminho)


if __name__ == "__main__":
    main()
